# -*- coding: utf-8 -*-
"""安全加固回归②：Excel 公式注入、登录 open redirect、认领失败残留照片。"""
import io
import os

import config

from conftest import db_conn, seed_item


def test_export_sanitizes_formula_values(admin_client):
    """以 = + - @ 开头的导出值必须被中和，Excel 打开不能当公式执行。"""
    from openpyxl import load_workbook
    seed_item(
        admin_client,
        name='=HYPERLINK("http://evil.example","点我")',
        description='-2+3',
        founder='@scanned',
    )
    resp = admin_client.get("/export")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    name_cell = ws.cell(row=2, column=2)      # 物品名称
    desc_cell = ws.cell(row=2, column=4)      # 特征描述
    founder_cell = ws.cell(row=2, column=8)   # 捡到人
    # 关键断言：任何单元格都不得是公式类型
    for cell in (name_cell, desc_cell, founder_cell):
        assert cell.data_type != "f", f"{cell.coordinate} 仍是公式: {cell.value!r}"
    # 中和方式：前导单引号，内容仍以文本形式可见
    assert name_cell.value.startswith("'")
    assert "=HYPERLINK" in name_cell.value


def test_export_normal_values_untouched(admin_client):
    """普通值不能被打扰（单引号前缀只加给危险字符开头的）。"""
    from openpyxl import load_workbook
    seed_item(admin_client, name="黑色钱包", founder="张护士")
    resp = admin_client.get("/export")
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    assert ws.cell(row=2, column=2).value == "黑色钱包"
    assert ws.cell(row=2, column=8).value == "张护士"


def test_login_next_url_rejects_protocol_relative(client):
    """session 里被塞入 //evil.com 形式的 next_url，登录成功后绝不能带偏。"""
    with client.session_transaction() as sess:
        sess["next_url"] = "//evil.example/steal"
    resp = client.post(
        "/login",
        data={"username": "liumin", "password": "liumin123"},
        follow_redirects=False,
    )
    assert resp.status_code == 302
    loc = resp.headers["Location"]
    assert not loc.startswith("//"), f"open redirect: {loc}"


def test_login_next_url_rejects_absolute_url(client):
    with client.session_transaction() as sess:
        sess["next_url"] = "https://evil.example/steal"
    resp = client.post(
        "/login",
        data={"username": "liumin", "password": "liumin123"},
        follow_redirects=False,
    )
    loc = resp.headers["Location"]
    assert loc.startswith("/") and not loc.startswith("//"), f"open redirect: {loc}"


def test_login_next_url_allows_local_path(client):
    """正常的站内路径不受影响（登录后跳回原页面）。"""
    with client.session_transaction() as sess:
        sess["next_url"] = "/list"
    resp = client.post(
        "/login",
        data={"username": "liumin", "password": "liumin123"},
        follow_redirects=False,
    )
    assert resp.headers["Location"] == "/list"


def test_failed_claim_cleans_orphaned_claimer_photo(admin_client):
    """认领竞争失败（两人都通过待认领检查、只有一人 UPDATE 命中）时，
    失败者上传的认领人照片不能残留在磁盘。"""
    import threading

    import app as app_mod

    admin_client.post(
        "/register",
        data={"name": "孤儿照片测试", "category": "其他", "storage_location": "抽屉"},
        headers={"X-Requested-With": "fetch"},
    )
    conn = db_conn()
    item_id = conn.execute("SELECT id FROM items LIMIT 1").fetchone()["id"]
    conn.close()

    png_a = b"\x89PNG\r\n\x1a\n" + b"\xAA" * 32
    png_b = b"\x89PNG\r\n\x1a\n" + b"\xBB" * 32

    def worker(photo, name):
        client = app_mod.app.test_client()
        client.post("/login", data={"username": "liumin", "password": "liumin123"})
        return client.post(
            "/claim",
            data={
                "item_id": item_id,
                "claimer_name": name,
                "feature_verified": "on",
                "claimer_photo_file": (io.BytesIO(photo), f"{name}.png"),
            },
            content_type="multipart/form-data",
            headers={"X-Requested-With": "fetch"},
        ).get_json()

    results = []
    threads = [
        threading.Thread(target=lambda: results.append(worker(png_a, "甲"))),
        threading.Thread(target=lambda: results.append(worker(png_b, "乙"))),
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert len(results) == 2
    assert sum(1 for r in results if r["ok"]) == 1, f"竞争结果异常: {results}"
    # 磁盘上最多只剩胜者的认领人照片，失败者的不许残留
    leftovers = [f for f in os.listdir(config.UPLOAD_FOLDER) if f.startswith("claimer_")]
    assert len(leftovers) <= 1, f"残留认领人照片: {leftovers}"
